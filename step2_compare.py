"""
Step 2d - Compare the alpha, heuristic and inductive miner models.

Reads the metrics written by the three miner programs (output/step2/*.json),
links every process variant to the Step 1 error types of its cases, and writes:
    output/step2/comparison.csv
    output/step2/variants.csv
    output/Step2_Process_Discovery.docx

Usage (after running the three step2_*_miner.py programs):
    python step2_compare.py
"""
import json

import pandas as pd
import pm4py
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

from pm_common import OUT, load_log
from step1_report_docx import TYPES, add_table, classify

MINERS = ["alpha", "heuristic", "inductive"]
VERSIONS = ["raw", "clean"]
METRIC_NAMES = {
    "cases": "Cases", "events": "Events", "variants": "Variants",
    "places": "Places", "transitions": "Transitions", "silent_transitions": "Silent (tau) transitions",
    "arcs": "Arcs", "fitness_tbr_perc_fit_traces": "Perfectly fitting traces (%)",
    "fitness_tbr_avg": "Fitness (token replay)", "fitness_alignments": "Fitness (alignments)",
    "precision": "Precision", "generalization": "Generalization", "simplicity": "Simplicity",
    "sound": "Sound (workflow net)",
}

# ---------------------------------------------------------------- metrics table
metrics = {(m, v): json.loads((OUT / f"{m}_{v}.json").read_text()) for m in MINERS for v in VERSIONS}
comp = pd.DataFrame({f"{m} ({v})": metrics[(m, v)] for v in VERSIONS for m in MINERS})
comp.index = [METRIC_NAMES[i] for i in comp.index]
comp.to_csv(OUT / "comparison.csv")
print(comp.to_string())

# ---------------------------------------------------------------- variants vs Step 1 errors
err = pd.read_csv("output/step1_errors.csv", dtype=str, keep_default_na=False)
ev_type = err.groupby("row").agg(case=("case", "first"), ts_raw=("ts_raw", "first"),
                                 checks=("check", lambda s: ", ".join(sorted(set(s)))))
ev_type["type"] = [classify(c, t) for c, t in zip(ev_type["checks"], ev_type["ts_raw"])]
case_types = ev_type.groupby("case")["type"].apply(lambda s: "+".join(sorted(set(s))))

log = load_log("raw")
traces = log.sort_values("time:timestamp").groupby("case:concept:name")["concept:name"].apply(" -> ".join)
var = traces.reset_index().rename(columns={"case:concept:name": "case", "concept:name": "variant"})
var["error_type"] = var["case"].map(case_types).fillna("none")
var_summary = (var.groupby("variant")
                  .agg(cases=("case", "count"),
                       error_types=("error_type", lambda s: ", ".join(f"{k}: {n}" for k, n in s.value_counts().items())),
                       example_cases=("case", lambda s: ", ".join(s.head(6))))
                  .sort_values("cases", ascending=False).reset_index())
var_summary.to_csv(OUT / "variants.csv", index=False)
print("\n" + var_summary.to_string())

# ---------------------------------------------------------------- Word report
COMMENTS = {
    "alpha": (
        "The alpha miner only looks at directly-follows pairs and has no noise handling. "
        "Because Approve Claim follows Evaluate Claim in most cases but precedes it in the 12 type C/D cases, "
        "it sees both orders and concludes the two activities are parallel. The missing timestamps "
        "(type A) add direct start->Verify Documents and Register->Evaluate connections, leaving tokens "
        "behind. The result is not a sound workflow net and no raw trace fits perfectly. Its token-replay "
        "precision of 1.0 is not reliable here, because replay on an unsound net with remaining tokens "
        "overstates precision."),
    "heuristic": (
        "The heuristic miner uses frequencies, so the frequent order Evaluate -> Approve keeps a strong "
        "dependency and the rare reversed order is mostly filtered out. It adds a silent transition to skip "
        "Register Claim (caused by the missing timestamp in Case_77) and an extra Evaluate -> Close place. "
        "The net is sound and fits most traces (84% perfectly), but the swapped/shifted cases do not replay."),
    "inductive": (
        "The inductive miner (noise threshold 0) must reproduce every trace, so it guarantees a sound net "
        "with fitness 1.0. To do so it makes Register and Verify skippable, makes Evaluate optional and puts "
        "Approve and Evaluate in parallel. This over-generalises: the model now allows approving a claim "
        "without evaluating it, which is why precision drops to 0.73. Every timestamp error is 'absorbed' "
        "into the model as legal behaviour."),
}

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Cm(1.5))
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10)

doc.add_heading("Step 2 - Process Discovery: Alpha, Heuristic and Inductive Miner", level=0)
doc.add_paragraph(
    "Each miner was run with PM4PY in its own program (step2_alpha_miner.py, step2_heuristic_miner.py, "
    "step2_inductive_miner.py) on two versions of the log: RAW, the log as given (events without a "
    "timestamp are dropped), and CLEAN, the log without the 22 cases flagged in Step 1. The clean log is "
    "not a correction (that is Step 5); it is used as a reference to show the effect of the errors.")

doc.add_heading("1. Process variants and the timestamp errors that cause them", level=1)
rows = [[i + 1, r["variant"], r["cases"], r["error_types"], r["example_cases"]] for i, r in var_summary.iterrows()]
add_table(doc, ["#", "Variant (raw log, ordered by timestamp)", "Cases", "Step 1 error types in these cases",
                "Example cases"], rows, [0.8, 11.0, 1.4, 5.5, 8.0])
legend = "; ".join(f"{k} = {v[0]}" for k, v in TYPES.items())
doc.add_paragraph(f"Error types: {legend}. Type B (truncated) does not change the order, so it creates no new variant.")

doc.add_heading("2. Metrics comparison", level=1)
header = ["Metric"] + list(comp.columns)
rows = [[name] + ["-" if pd.isna(x) else x for x in comp.loc[name]] for name in comp.index]
add_table(doc, header, rows, [6.2] + [3.4] * 6)
doc.add_paragraph(
    "Fitness: how much of the log the model can replay. Precision: how little extra behaviour the model "
    "allows. Generalization: how well the model avoids overfitting to the observed traces. Simplicity: "
    "based on the arc degree of the net (1 = simplest). Alignment fitness is '-' when the net is not sound.")

for m in MINERS:
    doc.add_page_break()
    doc.add_heading(f"3.{MINERS.index(m) + 1} {m.capitalize()} miner", level=1)
    doc.add_paragraph("Raw log:")
    doc.add_picture(str(OUT / f"{m}_raw.png"), width=Cm(25))
    doc.add_paragraph("Clean log (reference):")
    doc.add_picture(str(OUT / f"{m}_clean.png"), width=Cm(25))
    doc.add_paragraph(COMMENTS[m])

doc.add_heading("4. Conclusion", level=1)
for text in [
    "On the clean log all three miners discover the same sequential model "
    "Register -> Verify -> Evaluate -> Approve -> Close, with fitness, precision and simplicity equal to 1. "
    "So the differences between the models are caused only by the timestamp errors.",
    "Alpha miner: most sensitive. One wrong order turns a sequence into concurrency and the net becomes unsound.",
    "Heuristic miner: most robust. Frequency thresholds keep the main flow and treat errors as noise, at the "
    "cost of not fitting the affected traces.",
    "Inductive miner: always sound and fitting, but it over-generalises by making the errors part of the "
    "model (lower precision).",
    "Conclusion: the timestamps must be corrected (Step 5) before the discovered model can be trusted.",
]:
    doc.add_paragraph(text, style="List Bullet")

out = "output/Step2_Process_Discovery.docx"
doc.save(out)
print(f"\nSaved {out}")
